"""
generate_report.py
===================
Builds the Word document that explains the crowded predator-prey model, part by
part, with full equations, per-term derivations, literature sources, and the
honest access-confidence caveats from the five-paper-cluster literature pull.

Run AFTER crowded_model.py (it consumes the two PNG figures that script saves).
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
AMBER = RGBColor(0xC4, 0x78, 0x28)
GREY = RGBColor(0x55, 0x55, 0x55)
MONO = "Consolas"
SERIF = "Georgia"

doc = Document()

# ----- base styles -------------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = SERIF
normal.font.size = Pt(10.5)


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY if level <= 2 else AMBER
        run.font.name = SERIF
    return h


def body(text, italic=False, size=10.5, color=None, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = SERIF
    if color is not None:
        r.font.color.rgb = color
    return p


def equation(text):
    """Monospace, indented equation block."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = MONO
    r.font.size = Pt(9.5)
    r.font.color.rgb = NAVY
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.name = SERIF
        r.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.name = SERIF
    r2.font.size = Pt(10.5)
    return p


def caveat(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run("⚠  " + text)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = AMBER
    return p


def make_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = SERIF
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(str(val))
            run.font.size = Pt(8.5)
            run.font.name = SERIF
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Inches(w)
    return t


# ==============================================================================
# TITLE
# ==============================================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Modelling the Crowded Predator–Prey Oscillator")
r.bold = True
r.font.size = Pt(20)
r.font.name = SERIF
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("A bottom-up, parameter-renormalization model of macromolecular "
                "crowding (PEG 8000 vs Ficoll 400) in the Fujii–Rondelez "
                "DNA predator–prey PEN circuit")
r.italic = True
r.font.size = Pt(12)
r.font.name = SERIF
r.font.color.rgb = GREY

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("BIOMOD dry-lab · Step 3 (crowding layer) · model design "
                 "and implementation")
r.font.size = Pt(10)
r.font.color.rgb = GREY
doc.add_paragraph()

# ==============================================================================
# 1. EXECUTIVE SUMMARY
# ==============================================================================
heading("1.  Executive summary", 1)
body("This document specifies a mechanistic model of how molecular crowding "
     "reshapes the dynamics of our DNA predator–prey oscillator, and the "
     "Python that implements it (crowded_model.py). It is built directly on the "
     "47-paper Undermind literature scan and on a targeted re-reading of the "
     "~20 most load-bearing of those papers.")

body("The central finding of the literature is unambiguous: no published work "
     "models crowding in the Fujii–Rondelez oscillator directly. The only "
     "viable route is therefore to take the validated dilute kinetic backbone "
     "and let crowding renormalize its rate constants. We do exactly that, and "
     "we do not touch the equation structure — in particular the saturable-"
     "exonuclease term that is the verified oscillation engine stays exactly as "
     "validated in Step 1.")

body("Three design commitments follow from the literature and define the model:")
bullet("Crowding is decomposed into physically distinct channels (excluded "
       "volume, DNA hybridization thermodynamics, enzyme-class-specific "
       "catalysis, scale-dependent viscosity, and ion sequestration), each tied "
       "to a specific paper and each acting on the parameter it physically "
       "controls.", "Crowding is not one number.  ")
bullet("the renormalized rate constants are written as functions of the "
       "physical intermediates the report recommends fitting — volume "
       "fraction, correlation length, viscosity, water/ion activity — not as "
       "a direct empirical fit of the period. This is what makes the model "
       "predictive across both concentration and polymer molecular weight.",
       "Fit the intermediates, not the period.  ")
bullet("PEG 8000 (a flexible coil) and Ficoll 400 (a compact quasi-sphere) are "
       "given genuinely different physics — mesh-and-depletion vs hard-"
       "sphere excluded volume — rather than just different numbers, which "
       "is the strongest form of the project's novelty claim.",
       "Identity matters as much as concentration.  ")

body("Headline prediction (default coefficients). Crowding first lengthens the "
     "oscillation period — from 92 min in the dilute limit to ~170 min at 5% "
     "w/v PEG 8000 — because the exonuclease that drives the relaxation cycle "
     "is inhibited, then abolishes oscillation entirely above ~5% w/v as the "
     "system crosses the Hopf bifurcation into a stable fixed point. This is a "
     "concrete, falsifiable number for the wet lab to test, and it is "
     "consistent with the adjacent literature on damping of catalytic chemical "
     "oscillators under crowding.")
caveat("Every coefficient in the model is a literature-anchored, tunable "
       "hypothesis to be fitted to wet-lab calibration runs, not a measured "
       "constant for this exact system. Where a value is an extrapolation "
       "(most Ficoll-400 numbers; the Mg2+ ion channel) it is flagged as such.")

# ==============================================================================
# 2. THE DILUTE BASELINE
# ==============================================================================
heading("2.  The dilute baseline we build on", 1)
body("The validated Step-1 model is the Fujii Supporting-Information system "
     "(SI eqs 3–4) with the SI's measured PP1 parameters (Table S5). With N "
     "= prey and P = predator (nM), pol = Bst polymerase, rec = ttRecJ "
     "exonuclease:")
equation("dN/dt = k1·pol·G·N/(1 + b·G·N)  −  k2·pol·N·P  −  rec·kN·N/(1 + P/Km,P)")
equation("dP/dt = k2·pol·N·P  −  rec·kP·P/(1 + P/Km,P)")
body("The three terms are prey growth (saturating replication of N on template "
     "G by polymerase + nickase), predation (linear mass-action copying of the "
     "predator on the prey), and a shared saturable exonuclease sink. The paper "
     "proves (SI Fig S3 vs S5) that with first-order degradation the system can "
     "only damp to a fixed point; it is the saturable exonuclease — ttRecJ as "
     "a limited catalytic resource the predator can swamp — that creates the "
     "sustained limit cycle. That single fact dictates the whole crowding model: "
     "the exonuclease terms are the most dynamically sensitive place crowding "
     "can act.")

heading("Crowding enters by renormalization only", 2)
body("We replace each dilute rate constant by itself times a dimensionless "
     "crowding factor that depends on crowder concentration c, molecular weight "
     "M, and chemical identity:")
equation("k_i (c, M, identity)  =  k_i^0  ×  F_i(c, M, identity),    "
         "F_i → 1 as c → 0")
body("The structure of the ODEs is untouched; only the numbers k_i move. Each "
     "F_i is a product of channel factors, and the model is organized as three "
     "layers: control variables (what you set on the bench) → physical "
     "intermediates (Layer 1) → renormalized rate constants (Layer 2) → the "
     "unchanged ODEs run to give period, amplitude, and the bifurcation "
     "boundary (Layer 3).")

# ==============================================================================
# 3. ARCHITECTURE
# ==============================================================================
heading("3.  Model architecture (three layers)", 1)
make_table(
    ["Layer", "Contents", "Role"],
    [["0  Control", "concentration c (% w/v); molecular weight M; identity "
      "(PEG 8000 / Ficoll 400)", "what the experimentalist sets"],
     ["1  Physical intermediates", "volume fraction φ; overlap conc. c*; "
      "correlation length ξ; scale-dependent viscosity η_eff(r); free-ion "
      "fraction; identity geometry", "translate the controls into the "
      "quantities the literature actually measures"],
     ["2  Renormalized constants", "k1, k2, kN, kP, Km,P, b each = baseline "
      "× product of channel factors", "where crowding physically acts"],
     ["3  Dynamics / prediction", "the unchanged SI eqs 3–4, integrated", "period, "
      "amplitude, Hopf on/off vs (c, M, identity)"]],
    widths=[1.4, 3.2, 2.0])
body("")
body("The factor on each rate constant is assembled in log-space: each physical "
     "channel contributes an additive term to ln F_i, weighted by how strongly "
     "that reaction step couples to that channel. This keeps every contribution "
     "separable, signed, and individually traceable to a source:")
equation("ln F_i  =  Σ_channels  w[i][channel] · (channel log-contribution at φ)")

# ==============================================================================
# 4. LAYER 1 - PHYSICAL INTERMEDIATES
# ==============================================================================
heading("4.  Layer 1 — the physical intermediates", 1)

heading("4.1  Volume fraction φ", 2)
equation("φ = v̄ · c        (v̄ = partial specific volume; c in g/mL)")
body("The fraction of solution volume occupied by crowder. v̄ ≈ 0.83 mL/g for "
     "PEG and ≈ 0.67 mL/g for Ficoll, so 10% w/v PEG 8000 gives φ ≈ 0.083. "
     "φ is the master variable for every excluded-volume effect (Minton; Zhou, "
     "Rivas & Minton 2008).")

heading("4.2  Overlap concentration c* and the dilute/semidilute crossover", 2)
equation("c* ≈ 3M / (4π N_A R_g³)")
body("Below c* the polymer coils are separate (dilute); above c* they "
     "interpenetrate (semidilute) and a mesh forms. For PEG 8000, R_g ≈ 3.4 nm "
     "gives c* ≈ 8% w/v, so our experimental range straddles the crossover — "
     "exactly where the physics is richest. Ficoll, being compact, behaves as a "
     "hard sphere and has no semidilute mesh regime.")

heading("4.3  Correlation length (mesh size) ξ — PEG only", 2)
equation("ξ = R_g · (c/c*)^(−0.75)        (de Gennes, good solvent)")
body("The mesh size of the entangled polymer network. It diverges as c → 0 "
     "(no mesh in dilute solution), equals R_g at c = c*, and shrinks below R_g "
     "in the semidilute regime. ξ sets the crossover between the nano-viscosity "
     "a small species feels and the macro-viscosity a large enzyme feels. The "
     "−0.75 exponent is the standard de Gennes good-solvent scaling (Hołyst "
     "2009). Ficoll has no mesh, so this length does not apply to it — a "
     "concrete structural difference between the two crowders.")

heading("4.4  Scale-dependent viscosity η_eff(r)", 2)
body("The viscosity a diffusing species feels depends on its own size relative "
     "to the mesh. Hołyst's stretched-exponential form (PEG / coils):")
equation("η_eff(r)/η₀ = exp[ b · (R_eff/ξ)^a ],   1/R_eff² = 1/R_h² + 1/r²")
body("R_h is the crowder hydrodynamic radius, r the probe radius. Small species "
     "(short ssDNA, r ≈ 1.2 nm) have R_eff ≪ ξ and feel near-solvent "
     "viscosity; large enzymes (Bst, ttRecJ, r ≈ 3 nm) feel much more. We use "
     "a ≈ 0.78 (semidilute) and b ≈ 1.61 (PEG family). For Ficoll (hard "
     "spheres) we instead use the Mooney law η/η₀ = exp[[η]φ/(1−φ/φ_max)] "
     "with intrinsic viscosity [η] ≈ 2.5 and φ_max ≈ 0.6.")
caveat("a and b are Hołyst-family values taken from later papers in that group "
       "(Sozański 2016); the original 2009 paper states only 'constants close to "
       "unity'. PEG 8000 was not directly fitted — the nearest data are PEG "
       "6000/12000. Treat a, b as calibratable.")

heading("4.5  Free-ion fraction (optional ion-sequestration channel)", 2)
body("Nonionic crowders sequester buffer cations onto their monomer units, "
     "lowering the free-ion concentration that screens the DNA backbone "
     "(Bielec 2021). Per-monomer complexation constant κ ≈ 0.49 /M:")
equation("[ion]_free/[ion]_0 = 1/(1 + κ·c_monomer),   "
         "K_hyb ∝ [ion]^2.51")
body("Because like-charged DNA strands do not associate at all without "
     "counter-ion screening (Kowalski 2022), sequestering ions slows "
     "hybridization — a destabilizing pressure that opposes the excluded-volume "
     "stabilization. This channel is OFF by default.")
caveat("Bielec is parameterized for Na+; our oscillator buffer is Mg2+-based. "
       "Porting κ and the 2.51 exponent to Mg2+ is an extrapolation, which is "
       "why the channel is disabled by default and exposed only for exploration.")

# ==============================================================================
# 5. LAYER 2 - RENORMALIZED PARAMETERS
# ==============================================================================
heading("5.  Layer 2 — the renormalized rate constants", 1)
body("Five physical channels feed the rate constants. Each is written as a "
     "log-contribution at volume fraction φ (φ/(1−φ) is the scaled-particle "
     "variable Q). Signs: (+) speeds the step up, (−) slows it.")

heading("5.1  The channels (each tied to one mechanism and source)", 2)
make_table(
    ["Channel", "Log-contribution at φ", "Sign", "Mechanism & source"],
    [["HYB", "s_hyb · φ", "+", "net DNA hybridization stabilization; PEG 8000 "
      "k_close +4×, ΔΔG −1.73 kcal/mol at 10% (Baltierra-Jasso 2015); large-"
      "PEG duplex m = +0.156 kcal/molal (Knowles 2011)"],
     ["EV", "α_EV · φ/(1−φ)", "+", "excluded-volume / depletion enhancement "
      "of association & binding; scaled-particle theory (Minton); depletion "
      "Boltzmann factor e^(βΔU) (Berezhkovskii–Szabó 2016)"],
     ["PROC", "ρ_proc · φ", "+", "polymerase + nickase processivity / net "
      "replication; ~2–4× at 4% PEG (Akabayov 2013); endonuclease/nickase "
      "enhanced (Sasaki 2007); Pol–DNA binding up (Zimmerman 1987)"],
     ["EXO", "−e_exo · φ", "−", "exonuclease (ttRecJ) catalytic inhibition; "
      "Exo I Vmax −5.5× at 20% PEG 8000, Km unchanged (Sasaki 2007)"],
     ["ION", "2.51·ln([ion]_free/[ion]_0)", "−", "ion-sequestration "
      "de-screening of hybridization (Bielec 2021; Kowalski 2022); OFF by "
      "default"]],
    widths=[0.8, 1.9, 0.5, 3.4])
body("")
body("Viscosity is handled separately because the probe size differs by step: "
     "k1/k2 feel the polymerase-scale viscosity (slowdown), and − following "
     "Sasaki 2007, who found the nuclease effect is on Vmax, not diffusion — "
     "the exonuclease terms are NOT coupled to viscosity.")

heading("5.2  How each constant is built (the coupling weights)", 2)
body("Each rate constant couples to the channels with a weight in {0, 0.5, 1} "
     "reflecting how much that physical step is governed by that channel:")
make_table(
    ["Constant (step)", "HYB", "EV", "PROC", "EXO", "viscosity", "Net direction"],
    [["k1  prey growth (Pol+Nic on G)", "0.5", "0.5", "1.0", "0", "pol (−)",
      "viscosity dips it at low c, EV+PROC lift it at high c"],
     ["k2  predation (Pol on N·P)", "1.0", "1.0", "0.5", "0", "pol (−)",
      "up (bimolecular association strongly depletion-favoured)"],
     ["kN  prey decay (ttRecJ)", "0", "0.2", "0", "1.0", "none",
      "down (exonuclease inhibition)"],
     ["kP  predator decay (ttRecJ)", "0", "0.2", "0", "1.0", "none",
      "down (exonuclease inhibition)"],
     ["Km,P  exonuclease constant", "0", "0", "0", "0", "none",
      "flat (Sasaki: crowding moves Vmax, not Km)"],
     ["b  template saturation", "0", "0", "0", "0", "none", "flat (negligible)"]],
    widths=[1.9, 0.4, 0.4, 0.5, 0.4, 0.7, 2.0])
body("")
body("The Km,P decision is worth stressing: the generic Michaelis–Menten-under-"
     "crowding result (Weilandt 2018) lowers apparent Km, but the enzyme-class-"
     "specific nuclease measurements (Sasaki 2007) show the exonuclease effect "
     "lands on Vmax with Km essentially unchanged. We follow the enzyme-specific "
     "evidence and hold Km,P flat, putting the entire exonuclease effect on "
     "kN/kP. Because Km,P sits inside the saturable term that is the oscillation "
     "engine, this choice materially changes the predicted dynamics, so it is "
     "called out explicitly as a modelling decision.")

# ==============================================================================
# 6. THE CROWDED ODE SYSTEM
# ==============================================================================
heading("6.  The full crowded model, written out", 1)
body("Substituting the renormalized constants into the unchanged structure:")
equation("dN/dt = k1·F1·pol·G·N/(1 + b·Fb·G·N) − k2·F2·pol·N·P "
         "− rec·kN·FN·N/(1 + P/(Km,P·FKm))")
equation("dP/dt = k2·F2·pol·N·P − rec·kP·FP·P/(1 + P/(Km,P·FKm))")
body("with each factor assembled from Section 5:")
equation("ln F1 = 0.5·HYB + 0.5·EV + 1.0·PROC + 1.0·ION − ln η_eff(r_pol)")
equation("ln F2 = 1.0·HYB + 1.0·EV + 0.5·PROC + 1.0·ION − ln η_eff(r_pol)")
equation("ln FN = ln FP = −e_exo·φ + 0.2·EV")
equation("FKm = 1,   Fb = 1")
body("At φ = 0 every channel contribution vanishes, every F = 1, and the system "
     "is identical to the validated dilute model — a hard check the code "
     "asserts on every run.")

# ==============================================================================
# 7. IDENTITY AXIS
# ==============================================================================
heading("7.  The identity axis — PEG 8000 vs Ficoll 400", 1)
body("Polymer identity is not a single scalar; PEG and Ficoll differ in two "
     "physically separate ways, and the literature shows the 'more ideal' label "
     "even flips depending on the interaction partner.")
bullet("PEG 8000 is a flexible linear coil. It forms a semidilute mesh (Section "
       "4.3), so it gets scale-dependent viscosity and depletion attraction. "
       "Toward DNA folding it behaves as a near-ideal entropic excluded-volume "
       "crowder (Sung & Nesbitt 2021); toward enzymes/hydrophobic substrates it "
       "has genuine attractive enthalpic interactions that suppress activity.",
       "Geometry.  ")
bullet("Ficoll 400 is a compact, highly branched quasi-sphere. It has no mesh "
       "and is modelled as a hard-sphere excluded-volume crowder via the Mooney "
       "viscosity law. Its effect on enzyme rates is markedly milder than an "
       "equal mass of extended polymer (Homchaudhuri 2006: Ficoll slowed a "
       "reaction only 1.3–2.3× vs >5× for large dextran), because its "
       "quasi-spherical shape is close to an ideal hard sphere.",
       "Hard-sphere limit.  ")
body("In the model this shows up as: Ficoll uses the Mooney viscosity instead of "
     "the de Gennes mesh, and is given systematically milder channel strengths "
     "(s_hyb, α_EV, ρ_proc, e_exo all smaller). The molecular-weight axis "
     "enters through R_g, c*, and ξ, so the same model predicts both "
     "concentration and chain-length dependence — the report's recommended "
     "benchmark.")
caveat("No paper in the surveyed set used Ficoll 400. Every Ficoll-400 "
       "coefficient is an explicit extrapolation from compact-crowder behaviour "
       "and must be calibrated against our own wet-lab runs before any Ficoll "
       "number is trusted quantitatively.")

# ==============================================================================
# 8. PREDICTIONS
# ==============================================================================
heading("8.  Predictions and the falsifiable benchmark", 1)
body("Running the model across a concentration grid for both crowders gives the "
     "two predictions the project is judged on: how the period moves, and "
     "whether crowding switches the oscillation off (a Hopf-bifurcation "
     "crossing).")

heading("8.1  Rate-constant renormalization vs crowding", 2)
fig1 = os.path.join(HERE, "crowding_factors_peg.png")
if os.path.exists(fig1):
    doc.add_picture(fig1, width=Inches(5.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
body("Figure 1. How each PEG 8000 rate constant moves with concentration. "
     "Predation k2 rises (excluded-volume-favoured bimolecular association); the "
     "exonuclease decays kN, kP fall (catalytic inhibition); prey growth k1 dips "
     "then recovers as the polymerase-scale viscosity slowdown gives way to the "
     "excluded-volume and processivity gains.", italic=True, size=9)

heading("8.2  Period and amplitude — the headline result", 2)
fig2 = os.path.join(HERE, "crowding_period_prediction.png")
if os.path.exists(fig2):
    doc.add_picture(fig2, width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
body("Figure 2. Predicted oscillation period (left) and predator amplitude "
     "(right) vs crowder concentration, PEG 8000 vs Ficoll 400. The period "
     "lengthens as the exonuclease is inhibited; the amplitude collapses to zero "
     "at the Hopf boundary, where oscillation ceases.", italic=True, size=9)

make_table(
    ["Condition", "Period (min)", "Predator amp (nM)", "Oscillating?"],
    [["Dilute baseline", "92", "172", "yes"],
     ["5% w/v PEG 8000", "~172", "~178", "yes (period nearly doubled)"],
     ["10% w/v PEG 8000", "—", "0", "no (past Hopf, fixed point)"],
     ["PEG 8000 window", "—", "—", "oscillates ~0–5% w/v"],
     ["Ficoll 400 window", "—", "—", "oscillates ~0–4% w/v"]],
    widths=[2.0, 1.2, 1.6, 1.8])
body("")
body("The benchmark the report asks for — predicting period across "
     "concentration and molecular weight from physical intermediates rather than "
     "fitting the period directly — is met: period is an emergent output of the "
     "renormalized constants, and the model extends to MW through R_g, c*, and "
     "ξ. The single most testable claim is the critical crowder concentration "
     "(~5% w/v PEG 8000 with default coefficients) at which oscillation is lost.")

# ==============================================================================
# 9. ASSUMPTIONS & CALIBRATION
# ==============================================================================
heading("9.  Assumptions, limitations, and what needs calibration", 1)
bullet("The reduced two-variable model is assumed to remain valid under "
       "crowding (the SI's collapse of the full reactome to eqs 3–4 does not "
       "itself change; only its parameters do).", "Structure preserved.  ")
bullet("the period and Hopf boundary depend on a handful of channel strengths "
       "(s_hyb, α_EV, ρ_proc, e_exo and the Hołyst a, b). Defaults are "
       "anchored to specific papers but must be fitted to our enzymes and "
       "buffer.", "Coefficients are hypotheses.  ")
bullet("the strong exonuclease anchor (5.5×) is an Exo I number, not a "
       "measured ttRecJ number; it is a class-level prior. Calibrating ttRecJ "
       "activity vs PEG is the single highest-value wet-lab measurement.",
       "ttRecJ ≠ Exo I.  ")
bullet("the Bielec ion channel is Na+-based; the Fujii buffer is Mg2+. The "
       "channel is off by default and flagged where used.", "Mg2+ extrapolation.  ")
bullet("all Ficoll-400 numbers are extrapolated from compact-crowder behaviour; "
       "no primary Ficoll-400 datum exists in the surveyed literature.",
       "Ficoll uncalibrated.  ")
bullet("crowding-induced anomalous (fractal) kinetics (Schnell & Turner 2004) "
       "would make the rate coefficients time-dependent and break the autonomous-"
       "ODE form; we deliberately stay with concentration-dependent but time-"
       "constant rate constants, and note fractal kinetics as a possible "
       "advanced extension.", "Fractal kinetics not used.  ")

body("Priority calibration measurements (in order): (1) ttRecJ activity vs PEG "
     "8000 concentration; (2) prey/predator hybridization on-rate vs PEG "
     "concentration and MW; (3) solution viscosity vs PEG concentration and MW "
     "at 46.5 °C; (4) the same series for Ficoll 400.")

# ==============================================================================
# 10. REFERENCES
# ==============================================================================
heading("10.  References (cited subset of the 47-paper scan)", 1)
body("Access confidence is noted: [full] = full text read; [abs] = abstract "
     "only; [2°] = value taken from a later citing paper, not the original. "
     "Catalog numbers refer to the Undermind report.", italic=True, size=9)

refs = [
    "[#1] Fujii & Rondelez (2013) Predator-prey molecular ecosystems. ACS Nano "
    "7(1):27–34. — baseline PP model. [full, our Step-1 source]",
    "[#3] Montagne et al. (2011) Programming an in vitro DNA oscillator. Mol "
    "Syst Biol 7:466. — PEN oscillator backbone.",
    "[#4] Bielec et al. (2021) Ion complexation explains orders-of-magnitude "
    "changes in K. J Phys Chem Lett 13:112. — ION channel, κ≈0.49/M, "
    "K∝[Na+]^2.51. [full]",
    "[#6] Hou et al. (2011) Influence of nano-viscosity and depletion on DNA "
    "cleavage. Soft Matter 7:3092. — viscosity+depletion decomposition. [abs]",
    "[#7] Hołyst et al. (2009) Scaling form of viscosity at all length-scales in "
    "PEG. PCCP 11:9025. — scale-dependent viscosity η_eff/η₀=exp[b(R/ξ)^a]. "
    "[abs; a,b via 2° Sozański 2016]",
    "[#8] Knowles et al. (2011) Separation of preferential interaction and "
    "excluded-volume effects on DNA stability. PNAS 108:12699. — m-values, "
    "EV vs hydration split. [full]",
    "[#12] Berezhkovskii & Szabó (2016) Theory of crowding effects on "
    "bimolecular reaction rates. J Phys Chem B 120:5998. — Collins–Kimball + "
    "depletion e^(βΔU). [full]",
    "[#14] Zhou, Rivas & Minton (2008) Macromolecular crowding and confinement. "
    "Annu Rev Biophys 37:375. — enhancement magnitudes, thermo-vs-kinetic "
    "competition. [full]",
    "[#15/#42] Minton (2001 JBC; 1998 Methods Enzymol) — scaled-particle "
    "activity coefficients, two opposing effects. [reconstructed via #14]",
    "[#17] Weilandt & Hatzimanikatis (2018) Crowding effects on the Michaelis–"
    "Menten mechanism. Biophys J 117:355. — Km↓, Vmax↓, exp(β(φ)). [full]",
    "[#18] Baltierra-Jasso et al. (2015) Crowding-induced hybridization of "
    "single DNA hairpins. JACS 137:16020. — PEG 8000 k_close +4×, ΔΔG "
    "−1.73 kcal/mol. [full via proxy]",
    "[#19] Markarian & Schlenoff (2010) Crowding and ionic strength on "
    "oligonucleotide hybridization. J Phys Chem B 114:10620. — regime-"
    "dependent stabilization. [full]",
    "[#20] Sung & Nesbitt (2021) True crowding vs enthalpic interactions. J Phys "
    "Chem B 125:13147. — PEG entropic / Ficoll-dextran enthalpic toward DNA. "
    "[abs]",
    "[#21] Ghosh et al. (2020) Nearest-neighbor parameters under crowding. PNAS "
    "117:14194. — water-activity NN correction (PEG 200, destabilizing). [full]",
    "[#22] Sasaki, Miyoshi & Sugimoto (2006) Crowding and DNA polymerase. "
    "Biotechnol J 1:440. — Pol binding up, intrinsic kcat down. [abs]",
    "[#23] Sasaki, Miyoshi & Sugimoto (2007) Regulation of DNA nucleases by "
    "crowding. NAR 35:4086. — endo up, exo down ~5.5× (Vmax) at 20% PEG 8000. "
    "[full]",
    "[#24] Zimmerman & Harrison (1987) Crowding increases polymerase–DNA "
    "binding. PNAS 84:1871. — binding-up anchor. [abs]",
    "[#26] Kowalski et al. (2022) Screening of Coulomb repulsions accelerates "
    "like-charged reactions. Nat Commun 13:6451. — DNA association needs "
    "counter-ion screening. [full]",
    "[#27] Goins, Sanabria & Waxham (2008) Crowding and size effects on probe "
    "microviscosity. Biophys J 95:5362. — η_μ/η₀=exp(βC), β∝MW^0.11. [full]",
    "[#28] Banks & Fradin (2005) Anomalous diffusion of proteins due to "
    "crowding. Biophys J 89:2960. — subdiffusion α, limit ≈0.74. [full]",
    "[#34] Pastor et al. (2014) Crowding by dextrans in enzymatic reactions. "
    "Biophys Chem. — vmax & Km down, excluded-volume-tracked. [abs]",
    "[#35] Homchaudhuri et al. (2006) Crowding by dextrans and Ficolls, size-"
    "dependent. Biopolymers 83:477. — Ficoll milder, shape-dependent. [abs]",
    "[#38] Akabayov et al. (2013) Impact of crowding on DNA replication. Nat "
    "Commun 4:1615. — replication up ~2–4× at 4% PEG via compaction. [full]",
    "[#46] Echeverría et al. (2018) Damping and clustering of catalytic chemical "
    "oscillators in crowding. Physica A. — adjacent precedent for crowding-"
    "induced damping.",
]
for r_ in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    run = p.add_run(r_)
    run.font.size = Pt(8.5)
    run.font.name = SERIF

# footer note
doc.add_paragraph()
note = doc.add_paragraph()
r = note.add_run("Implementation: crowded_model.py (self-contained; reproduces "
                 "the validated dilute limit exactly, then renormalizes every "
                 "rate constant per the channels above). All coefficients are "
                 "tunable hypotheses for wet-lab calibration.")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = GREY

out = os.path.join(HERE, "Crowded_PredatorPrey_Model.docx")
doc.save(out)
print(f"saved {out}")
